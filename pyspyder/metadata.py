import csv
import logging
import os

logger = logging.getLogger("pyspyder")


def extract_metadata(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".pptx": _extract_pptx,
        ".doc": _extract_ole,
        ".xls": _extract_ole,
        ".ppt": _extract_ole,
    }

    extractor = extractors.get(ext)
    if not extractor:
        logger.debug(f"No metadata extractor for {ext} files")
        return {}

    try:
        return extractor(filepath)
    except Exception as err:
        logger.warning(f"Failed to extract metadata from {os.path.basename(filepath)}: {err}")
        return {}


def _extract_pdf(filepath):
    from PyPDF2 import PdfReader

    reader = PdfReader(filepath)
    info = reader.metadata
    if not info:
        return {}

    metadata = {}
    field_map = {
        "/Author": "Author",
        "/Creator": "Creator",
        "/Producer": "Producer",
        "/Title": "Title",
        "/Subject": "Subject",
        "/CreationDate": "Created",
        "/ModDate": "Modified",
    }
    for pdf_key, our_key in field_map.items():
        value = info.get(pdf_key)
        if value:
            metadata[our_key] = str(value).strip()

    return metadata


def _extract_docx(filepath):
    from docx import Document

    doc = Document(filepath)
    props = doc.core_properties
    metadata = {}

    if props.author:
        metadata["Author"] = props.author
    if props.last_modified_by:
        metadata["LastModifiedBy"] = props.last_modified_by
    if props.title:
        metadata["Title"] = props.title
    if props.subject:
        metadata["Subject"] = props.subject
    if props.created:
        metadata["Created"] = str(props.created)
    if props.modified:
        metadata["Modified"] = str(props.modified)

    return metadata


def _extract_xlsx(filepath):
    from openpyxl import load_workbook

    wb = load_workbook(filepath, read_only=True, data_only=True)
    props = wb.properties
    metadata = {}

    if props.creator:
        metadata["Author"] = props.creator
    if props.lastModifiedBy:
        metadata["LastModifiedBy"] = props.lastModifiedBy
    if props.title:
        metadata["Title"] = props.title
    if props.created:
        metadata["Created"] = str(props.created)
    if props.modified:
        metadata["Modified"] = str(props.modified)

    wb.close()
    return metadata


def _extract_pptx(filepath):
    from pptx import Presentation

    prs = Presentation(filepath)
    props = prs.core_properties
    metadata = {}

    if props.author:
        metadata["Author"] = props.author
    if props.last_modified_by:
        metadata["LastModifiedBy"] = props.last_modified_by
    if props.title:
        metadata["Title"] = props.title
    if props.subject:
        metadata["Subject"] = props.subject
    if props.created:
        metadata["Created"] = str(props.created)
    if props.modified:
        metadata["Modified"] = str(props.modified)

    return metadata


def _extract_ole(filepath):
    import olefile

    if not olefile.isOleFile(filepath):
        return {}

    ole = olefile.OleFileIO(filepath)
    metadata = {}

    try:
        meta = ole.get_metadata()

        if meta.author:
            metadata["Author"] = meta.author.decode("utf-8", errors="replace")
        if meta.last_saved_by:
            metadata["LastModifiedBy"] = meta.last_saved_by.decode("utf-8", errors="replace")
        if meta.company:
            metadata["Company"] = meta.company.decode("utf-8", errors="replace")
        if meta.title:
            metadata["Title"] = meta.title.decode("utf-8", errors="replace")
        if meta.subject:
            metadata["Subject"] = meta.subject.decode("utf-8", errors="replace")
        if meta.creating_application:
            metadata["CreatingApplication"] = meta.creating_application.decode("utf-8", errors="replace")
        if meta.create_time:
            metadata["Created"] = str(meta.create_time)
        if meta.last_saved_time:
            metadata["Modified"] = str(meta.last_saved_time)
    finally:
        ole.close()

    return metadata


def format_metadata_table(results):
    if not results:
        return "No metadata found."

    lines = []
    lines.append("")
    lines.append("=" * 80)
    lines.append("METADATA RESULTS")
    lines.append("=" * 80)

    for entry in results:
        filename = entry["file"]
        metadata = entry["metadata"]

        if not metadata:
            continue

        lines.append(f"\n  {filename}")
        lines.append(f"  {'-' * len(filename)}")
        for key, value in metadata.items():
            lines.append(f"    {key:20s}: {value}")

    lines.append("")
    lines.append("=" * 80)

    unique_users = set()
    unique_software = set()
    for entry in results:
        meta = entry["metadata"]
        for field in ["Author", "LastModifiedBy"]:
            if field in meta and meta[field]:
                unique_users.add(meta[field])
        for field in ["Creator", "Producer", "CreatingApplication"]:
            if field in meta and meta[field]:
                unique_software.add(meta[field])

    if unique_users:
        lines.append("UNIQUE USERS FOUND:")
        for user in sorted(unique_users):
            lines.append(f"  - {user}")
        lines.append("")

    if unique_software:
        lines.append("SOFTWARE IDENTIFIED:")
        for sw in sorted(unique_software):
            lines.append(f"  - {sw}")

    if unique_users or unique_software:
        lines.append("=" * 80)

    return "\n".join(lines)


def export_csv(results, csv_path):
    if not results:
        logger.warning("No results to export.")
        return

    all_keys = set()
    for entry in results:
        all_keys.update(entry["metadata"].keys())
    all_keys = sorted(all_keys)

    fieldnames = ["Filename", "URL"] + all_keys

    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()

        for entry in results:
            row = {"Filename": entry["file"], "URL": entry.get("url", "")}
            row.update(entry["metadata"])
            writer.writerow(row)

    logger.info(f"Metadata exported to {csv_path}")
