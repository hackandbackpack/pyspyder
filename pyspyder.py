#!/usr/bin/env python3
"""PySpyder - Spider a target domain for files, download them, and extract metadata."""

import argparse
import csv
import logging
import os
import re
import sys
import time
import xml.etree.ElementTree as ElementTree
from collections import deque
from datetime import datetime
from urllib.parse import parse_qsl, unquote, urlencode, urljoin, urlparse, urlsplit
from urllib.robotparser import RobotFileParser

REQUIRED_PACKAGES = {
    "requests": "requests",
    "bs4": "beautifulsoup4",
    "lxml": "lxml",
    "pypdf": "pypdf",
    "docx": "python-docx",
    "openpyxl": "openpyxl",
    "pptx": "python-pptx",
    "olefile": "olefile",
}


def check_dependencies():
    missing = []
    for module, package in REQUIRED_PACKAGES.items():
        try:
            __import__(module)
        except ImportError:
            missing.append(package)

    if missing:
        print(f"Error: Missing required packages: {', '.join(missing)}")
        print(f"Install them with: pip install {' '.join(missing)}")
        sys.exit(1)


check_dependencies()

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter, Retry

__version__ = "1.1.0"

logger = logging.getLogger("pyspyder")


# Recent desktop Firefox. A plain Chrome UA is a common WAF/security-plugin
# block signal - a stock WordPress security rule will 403 the Chrome token while
# letting Firefox through - and a silent 403 on the seed URL reads as "0 files
# found". Override with --user-agent when a target wants something else. Bump the
# version now and then so it stays current.
DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:133.0) "
    "Gecko/20100101 Firefox/133.0"
)

DEFAULT_FILE_TYPES = ["pdf", "docx", "xlsx", "doc", "xls", "pptx", "ppt"]

EMAIL_REGEX = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")

# Metadata is authored by whoever made the document, which on an engagement means
# the target. It reaches an operator's terminal and their spreadsheet, so strip
# anything that could reprogram a terminal and neutralize anything a spreadsheet
# would run as a formula.
CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b-\x1f\x7f-\x9f]")
CSV_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
MAX_METADATA_LENGTH = 500

MAX_FILENAME_LENGTH = 120
MAX_EXTENSION_LENGTH = 12
MAX_APP_XML_BYTES = 1024 * 1024
MAX_SITEMAPS = 50

# Query parameters that never change which page you get back. Dropping them stops
# the crawler from fetching the same page once per campaign tag.
TRACKING_PARAMS = {
    "utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content",
    "utm_id", "gclid", "fbclid", "msclkid", "mc_cid", "mc_eid", "_ga",
}

# Names Windows treats as devices rather than files, with or without an extension.
WINDOWS_RESERVED_NAMES = (
    {"CON", "PRN", "AUX", "NUL"}
    | {f"COM{i}" for i in range(1, 10)}
    | {f"LPT{i}" for i in range(1, 10)}
)


# ---------------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------------

def setup_logging(verbose=False):
    level = logging.DEBUG if verbose else logging.INFO
    logging.basicConfig(level=level, format="%(levelname)s: %(message)s")
    # urllib3 announces every connection retry at WARNING, and pypdf complains
    # about every malformed document. Both bury our own output, and we already
    # report the outcome of each fetch and each extraction ourselves.
    if not verbose:
        logging.getLogger("urllib3").setLevel(logging.ERROR)
        logging.getLogger("pypdf").setLevel(logging.ERROR)
    return logging.getLogger("pyspyder")


def get_session(user_agent=None):
    session = requests.Session()
    session.headers.update({
        "User-Agent": user_agent or DEFAULT_USER_AGENT,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
        # Only what requests can decode without extra packages. Advertising br
        # while no brotli decoder is installed handed us undecoded bytes on every
        # Brotli server - the page parsed to zero links and the crawl reported
        # nothing found. Add brotli to requirements before putting br back.
        "Accept-Encoding": "gzip, deflate",
        "DNT": "1",
        "Connection": "keep-alive",
        "Upgrade-Insecure-Requests": "1",
    })
    retry_strategy = Retry(
        total=3,
        backoff_factor=0.5,
        status_forcelist=[429, 500, 502, 503],
        # urllib3 sleeps a Retry-After header verbatim with no ceiling, so one
        # "Retry-After: 86400" would park the run for a day. Our own backoff is
        # capped at urllib3's backoff_max.
        respect_retry_after_header=False,
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


def normalize_domain(raw):
    """Turn whatever the user typed into a bare host[:port].

    Accepts 'example.com', 'https://example.com', 'EXAMPLE.com/docs/' and so on.
    Returns (host, discarded_path) so the caller can warn about a dropped path.
    """
    value = raw.strip()
    if "//" in value:
        value = value.split("//", 1)[1]
    host, _, remainder = value.partition("/")
    host = host.split("@")[-1].strip().rstrip(".").lower()
    return host, remainder.strip("/")


def resolve_target(domain, session):
    """Find the scheme and hostname that actually serve this domain.

    Returns (scheme, host). HTTPS is tried first, then plain HTTP, so hosts that
    only speak HTTP are still reachable. Some servers reject HEAD, so a rejected
    HEAD is retried as a GET before giving up on a scheme.
    """
    for scheme in ("https", "http"):
        for method in (session.head, session.get):
            try:
                # stream=True so a GET fallback doesn't pull a whole homepage
                # into memory just to read the final URL off it.
                resp = method(f"{scheme}://{domain}/", timeout=10,
                              allow_redirects=True, stream=True)
            except Exception:
                continue
            try:
                if resp.status_code >= 400:
                    continue
                parsed = urlparse(resp.url)
                host = parsed.netloc.lower()
                return (parsed.scheme.lower() or scheme), (host or domain)
            finally:
                resp.close()
    return "https", domain


def url_key(url):
    """Canonical key identifying one resource, for the visited set and dedup.

    This is deliberately lossy - it drops the scheme and rewrites the query - so
    it must never be requested. Fetches always use the URL as it was discovered,
    because a rewritten query is a different request as far as the server is
    concerned ('?flag' is not '?flag=', and a signed URL stops verifying).
    """
    parsed = urlparse(url)

    query = parsed.query
    if query:
        kept = [(k, v) for k, v in parse_qsl(query, keep_blank_values=True)
                if k.lower() not in TRACKING_PARAMS]
        query = urlencode(sorted(kept))

    # Taken straight off the netloc rather than via parsed.port, which raises
    # ValueError on a malformed port and would abort the whole crawl over one
    # bad link. An unfetchable URL just fails later like any other.
    host = parsed.netloc.rsplit("@", 1)[-1].lower().rstrip(".")
    key = f"{host}{parsed.path.rstrip('/')}"
    return f"{key}?{query}" if query else key


def bare_host(host_or_netloc):
    """Lowercase hostname with any port stripped. Handles IPv6 literals."""
    if not host_or_netloc:
        return ""
    try:
        return (urlsplit(f"//{host_or_netloc}").hostname or "").rstrip(".")
    except ValueError:
        return ""


def host_in_scope(link_host, scope_hosts):
    """True only for the target host itself or a subdomain of it.

    A substring test would accept example.com.attacker.tld and notexample.com,
    which is how a spider ends up crawling somebody else's estate.
    """
    host = bare_host(link_host)
    if not host:
        return False
    # Both sides go through bare_host, so a caller passing "[::1]:8443" or
    # "Example.com:443" as scope still compares correctly.
    for scope in scope_hosts:
        scope = bare_host(scope)
        if scope and (host == scope or host.endswith(f".{scope}")):
            return True
    return False


def normalize_file_types(raw):
    """Split a -f value into bare lowercase extensions.

    Accepts 'pdf,docx', '.pdf, .docx' and 'PDF' alike - a leading dot is the
    natural way to type an extension and used to match nothing at all.
    """
    return [ft.strip().lstrip(".").lower() for ft in raw.split(",") if ft.strip(" .")]


def is_target_file(url, file_types):
    parsed = urlparse(url)
    path_lower = parsed.path.lower()
    return any(path_lower.endswith(f".{ft.lower().lstrip('.')}") for ft in file_types)


def generate_output_dir(domain):
    timestamp = datetime.now().strftime("%Y-%m-%d-%H%M%S")
    safe_domain = re.sub(r"[^\w.-]", "_", domain)
    return f"{safe_domain}-{timestamp}"


def _progress(message):
    """Single-line crawl counter. Skipped when stdout is redirected."""
    if sys.stdout.isatty():
        print(f"\r{message}    ", end="", flush=True)


def _progress_clear():
    if sys.stdout.isatty():
        print("\r" + " " * 78 + "\r", end="", flush=True)


# ---------------------------------------------------------------------------
# Spider
# ---------------------------------------------------------------------------

def spider_domain(domain, file_types, session, max_depth=2, delay=1.0,
                  ignore_robots=False, max_pages=500):
    file_urls = set()
    visited = set()

    domain, _ = normalize_domain(domain)
    logger.info(f"Spider: Starting crawl of {domain} (depth={max_depth}, max_pages={max_pages})")

    scheme, actual_domain = resolve_target(domain, session)
    if actual_domain != domain:
        logger.info(f"Spider: {domain} redirects to {actual_domain}")
    if scheme != "https":
        logger.info(f"Spider: {actual_domain} does not serve HTTPS, using {scheme}://")

    base_url = f"{scheme}://{actual_domain}"
    queue = deque([(base_url + "/", 0)])
    scope_hosts = {bare_host(domain), bare_host(actual_domain)}

    robots = RobotsPolicy(session, ignore=ignore_robots)
    robots_sitemaps = robots.sitemaps_for(scheme, actual_domain)
    requested_delay = robots.crawl_delay(scheme, actual_domain)
    if requested_delay and requested_delay > delay:
        logger.info(f"Spider: robots.txt requests a {requested_delay}s crawl delay; "
                    f"running with --delay {delay}")

    while queue:
        if len(visited) >= max_pages:
            logger.info(f"Spider: Reached page limit ({max_pages}), stopping crawl")
            break

        current_url, depth = queue.popleft()
        key = url_key(current_url)

        if key in visited:
            continue
        visited.add(key)

        if not robots.allowed(current_url):
            continue

        _progress(f"  Crawling: {len(visited)} pages | {len(file_urls)} files found "
                  f"| depth {depth}/{max_depth}")
        logger.debug(f"Spider: Crawling {current_url} (depth {depth})")

        response, status = _fetch_page(session, current_url)
        # One delay covering every request, wherever the iteration ends up next.
        time.sleep(delay)
        if response is None:
            # depth 0 is the seed - if the root won't load, the whole run is
            # blind, so say so instead of quietly reporting nothing found.
            if depth == 0:
                _warn_blocked_seed(current_url, status)
            continue

        content_type = response.headers.get("Content-Type", "").lower()
        if not any(kind in content_type for kind in ("text/html", "application/xhtml")):
            continue

        # A redirect means the page in hand isn't the URL we asked for, and
        # relative links have to resolve against where we actually landed.
        final_url = response.url or current_url
        if url_key(final_url) != key:
            if not host_in_scope(urlparse(final_url).netloc, scope_hosts):
                logger.debug(f"Spider: Redirected out of scope, skipping {final_url}")
                continue
            visited.add(url_key(final_url))

        soup = BeautifulSoup(response.text, "html.parser")

        # <base href> overrides the document's own URL for relative links.
        base_tag = soup.find("base", href=True)
        link_base = urljoin(final_url, base_tag["href"]) if base_tag else final_url

        for anchor in soup.find_all("a", href=True):
            absolute_url = urljoin(link_base, anchor["href"])
            parsed = urlparse(absolute_url)

            if parsed.scheme not in ("http", "https"):
                continue
            if not host_in_scope(parsed.netloc, scope_hosts):
                logger.debug(f"Spider: Out of scope, skipping {absolute_url}")
                continue

            clean_url = parsed._replace(fragment="").geturl()

            if is_target_file(clean_url, file_types):
                if robots.allowed(clean_url):
                    file_urls.add(clean_url)
            elif depth < max_depth and url_key(clean_url) not in visited:
                queue.append((clean_url, depth + 1))

    _progress_clear()

    logger.info("Spider: Checking sitemap.xml for additional files...")
    file_urls.update(_parse_sitemap(base_url, file_types, session, scope_hosts,
                                    robots, delay, robots_sitemaps))

    if robots.blocked:
        logger.info(f"Spider: Skipped {robots.blocked} URL(s) disallowed by robots.txt "
                    f"(use --ignore-robots to collect them)")

    logger.info(f"Spider: Found {len(file_urls)} file(s) across {len(visited)} pages crawled")
    return file_urls


def _fetch_page(session, url):
    """GET one page. Returns (response, status).

    response is None when the page could not be used; status is the HTTP status
    the server returned (so the caller can tell a 403 block from a dead link),
    or None when the request never got an answer at all.
    """
    try:
        response = session.get(url, timeout=15)
    except Exception as err:
        logger.debug(f"Spider: Failed to fetch {url}: {err}")
        return None, None
    if response.status_code >= 400:
        logger.debug(f"Spider: {url} returned HTTP {response.status_code}")
        return None, response.status_code
    return response, None


def _warn_blocked_seed(url, status):
    """Loud warning when the starting page itself can't be fetched.

    A silent failure here reads as "0 files found" and the operator concludes the
    target exposes nothing, when in reality the crawl never got in the door.
    """
    if status in (401, 403, 429):
        logger.warning(
            f"Spider: the starting page {url} returned HTTP {status}. The site is "
            f"likely blocking the crawler (a WAF or User-Agent filter) - try a "
            f"different --user-agent. No files can be found until this is resolved."
        )
    elif status is not None:
        logger.warning(
            f"Spider: the starting page {url} returned HTTP {status}; there is "
            f"nothing to crawl."
        )
    else:
        logger.warning(
            f"Spider: could not reach the starting page {url}; check the domain "
            f"and network connectivity."
        )


class RobotsPolicy:
    """robots.txt rules, fetched once per host and cached.

    Rules are per host, so www.example.com saying Disallow: /docs/ must not
    suppress files served from cdn.example.com, which has its own robots.txt.
    """

    def __init__(self, session, ignore=False):
        self.session = session
        self.ignore = ignore
        self.parsers = {}
        self.sitemaps = {}
        self.blocked = 0

    def allowed(self, url):
        if self.ignore:
            return True
        parsed = urlparse(url)
        if not parsed.netloc:
            return True
        parser = self._parser(parsed.scheme or "https", parsed.netloc)
        if parser and not parser.can_fetch("*", url):
            self.blocked += 1
            logger.debug(f"Robots: Disallowed {url}")
            return False
        return True

    def sitemaps_for(self, scheme, netloc):
        """Sitemap: URLs advertised in that host's robots.txt."""
        self._parser(scheme, netloc)
        return list(self.sitemaps.get(netloc, []))

    def crawl_delay(self, scheme, netloc):
        parser = self._parser(scheme, netloc)
        return parser.crawl_delay("*") if parser else None

    def _parser(self, scheme, netloc):
        if netloc not in self.parsers:
            self.parsers[netloc] = self._load(scheme, netloc)
        return self.parsers[netloc]

    def _load(self, scheme, netloc):
        robots_url = f"{scheme}://{netloc}/robots.txt"
        self.sitemaps[netloc] = []
        try:
            response = self.session.get(robots_url, timeout=10)
            if response.status_code != 200:
                logger.debug(f"Robots: None at {robots_url} (HTTP {response.status_code})")
                return None
        except Exception as err:
            logger.debug(f"Robots: Could not fetch {robots_url}: {err}")
            return None

        lines = response.text.splitlines()
        for line in lines:
            name, _, value = line.partition(":")
            if name.strip().lower() == "sitemap" and value.strip():
                self.sitemaps[netloc].append(value.strip())

        parser = RobotFileParser()
        parser.parse(lines)
        logger.debug(f"Robots: Loaded {robots_url}")
        return parser


def _parse_sitemap(base_url, file_types, session, scope_hosts, robots, delay,
                   extra_sitemaps=None):
    """Parse sitemap.xml (and nested sitemaps) for in-scope file URLs."""
    file_urls = set()
    sitemap_urls = [f"{base_url}/sitemap.xml"]
    # Sitemap: lines are supposed to be absolute but plenty of sites write them
    # relative, and those would otherwise be dropped as having no host at all.
    sitemap_urls.extend(urljoin(f"{base_url}/", entry) for entry in extra_sitemaps or [])
    visited_sitemaps = set()

    while sitemap_urls and len(visited_sitemaps) < MAX_SITEMAPS:
        sitemap_url = sitemap_urls.pop()
        if sitemap_url in visited_sitemaps:
            continue
        visited_sitemaps.add(sitemap_url)

        # A Sitemap: line in robots.txt can name any host at all, so this is
        # checked here rather than only on nested <sitemap> entries.
        if not host_in_scope(urlparse(sitemap_url).netloc, scope_hosts):
            logger.debug(f"Sitemap: Out of scope, skipping {sitemap_url}")
            continue

        response = None
        try:
            response = session.get(sitemap_url, timeout=10)
        except Exception as err:
            logger.debug(f"Sitemap: Could not fetch {sitemap_url}: {err}")
        time.sleep(delay)
        if response is None or response.status_code != 200:
            continue

        try:
            soup = BeautifulSoup(response.text, "xml")
        except Exception as err:
            # A missing parser or unusable body must not sink the whole run -
            # the crawl results are already in hand by this point.
            logger.warning(f"Sitemap: Could not parse {sitemap_url}: {err}")
            continue

        for loc in soup.find_all("loc"):
            url = loc.get_text(strip=True)
            if not is_target_file(url, file_types):
                continue
            if not host_in_scope(urlparse(url).netloc, scope_hosts):
                logger.debug(f"Sitemap: Out of scope, skipping {url}")
                continue
            if robots.allowed(url):
                file_urls.add(url)

        for sitemap_tag in soup.find_all("sitemap"):
            loc = sitemap_tag.find("loc")
            if loc is None:
                continue
            nested_url = loc.get_text(strip=True)
            if nested_url not in visited_sitemaps:
                sitemap_urls.append(nested_url)

    if sitemap_urls:
        logger.warning(f"Sitemap: Stopped after {MAX_SITEMAPS} sitemaps, "
                       f"{len(sitemap_urls)} left unread")

    if file_urls:
        logger.info(f"Sitemap: Found {len(file_urls)} file(s) in sitemap(s)")
    else:
        logger.debug("Sitemap: No target files found in sitemap(s)")

    return file_urls


# ---------------------------------------------------------------------------
# Downloader
# ---------------------------------------------------------------------------

def download_files(urls, output_dir, session, delay=0.5, max_bytes=0, scope_hosts=None):
    os.makedirs(output_dir, exist_ok=True)

    downloaded = []
    skipped = 0
    used_names = set()
    urls = list(urls)
    total = len(urls)

    for idx, url in enumerate(urls, 1):
        filename = _unique_filename(_extract_filename(url), used_names)
        filepath = os.path.join(output_dir, filename)
        label = f"[{idx}/{total}]"

        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            logger.info(f"{label} Skipped (exists): {filename}")
            downloaded.append((url, filepath))
            skipped += 1
            continue  # nothing was requested, so no delay is owed

        if _fetch_to_file(url, filepath, session, max_bytes, label, scope_hosts):
            logger.info(f"{label} Downloaded: {filename}")
            downloaded.append((url, filepath))
        else:
            # Hand the name back so a later file can use it.
            used_names.discard(filename)

        time.sleep(delay)

    if skipped:
        logger.info(f"Skipped {skipped} file(s) that already existed on disk")

    return downloaded


def _fetch_to_file(url, filepath, session, max_bytes, label, scope_hosts=None):
    """Fetch one URL into filepath. True only when a complete file is in place.

    The body streams into a .part file that is renamed at the end, so an
    interrupted transfer never leaves something that looks like a whole file.
    """
    filename = os.path.basename(filepath)
    partial = filepath + ".part"
    try:
        response = session.get(url, timeout=30, stream=True)
        response.raise_for_status()

        # Documents are often redirected onto a CDN, which is legitimate and
        # worth following - but on a scoped engagement it should be on the
        # record that bytes came from a host the operator did not name.
        if scope_hosts:
            final_host = urlparse(response.url or url).netloc
            if final_host and not host_in_scope(final_host, scope_hosts):
                logger.info(f"{label} Note: {filename} redirected to off-scope host "
                            f"{bare_host(final_host)}")

        content_type = response.headers.get("Content-Type", "").lower()
        if _is_html_response(content_type, filename):
            logger.warning(f"{label} Skipped (server returned HTML): {filename}")
            return False

        declared = response.headers.get("Content-Length", "")
        if max_bytes and declared.isdigit() and int(declared) > max_bytes:
            logger.warning(f"{label} Skipped (larger than limit): {filename} "
                           f"({int(declared)} bytes)")
            return False

        written = 0
        with open(partial, "wb") as fh:
            for chunk in response.iter_content(chunk_size=8192):
                written += len(chunk)
                if max_bytes and written > max_bytes:
                    raise ValueError(f"exceeds --max-file-size at {written} bytes")
                fh.write(chunk)

        if not _validate_file_magic(partial):
            logger.warning(f"{label} Warning: {filename} may be corrupt "
                           f"(content doesn't match extension)")

        os.replace(partial, filepath)
        return True

    except Exception as err:
        logger.warning(f"{label} Failed to download {url}: {err}")
        return False
    finally:
        # Also runs on Ctrl-C, so a cancelled run leaves no stray .part files.
        if os.path.exists(partial):
            try:
                os.remove(partial)
            except OSError:
                logger.debug(f"Could not remove partial file {partial}")


def _unique_filename(filename, used_names):
    """Pick a name no other file in this run has taken."""
    if filename not in used_names:
        used_names.add(filename)
        return filename

    name, ext = os.path.splitext(filename)
    counter = 1
    while f"{name}_{counter}{ext}" in used_names:
        counter += 1
    unique = f"{name}_{counter}{ext}"
    used_names.add(unique)
    return unique


def _extract_filename(url):
    parsed = urlparse(url)
    path = unquote(parsed.path)
    filename = os.path.basename(path.replace("\\", "/"))
    filename = "".join(c for c in filename if c.isalnum() or c in "._- ")
    # Windows silently rejects trailing dots and spaces.
    filename = filename.strip().rstrip(". ")

    name, ext = os.path.splitext(filename)
    # A name of ".pdf" splits to name=".pdf", ext="" and would lose the very
    # extension the extractors dispatch on.
    if not ext and name.startswith("."):
        name, ext = "", name
    if name.upper() in WINDOWS_RESERVED_NAMES:
        name = f"_{name}"
    # Overlong names would only fail later as an opaque OSError, once the output
    # directory pushes the full path past the platform limit. The extension gets
    # trimmed too, since it can be the long part.
    if len(name) + len(ext) > MAX_FILENAME_LENGTH:
        ext = ext[:MAX_EXTENSION_LENGTH]
        name = name[:max(1, MAX_FILENAME_LENGTH - len(ext))].rstrip(". ")
    if not name:
        return f"unknown_file{ext}" if ext else "unknown_file"
    return f"{name}{ext}"


MAGIC_BYTES = {
    ".pdf": b"%PDF",
    ".doc": b"\xd0\xcf\x11\xe0",
    ".xls": b"\xd0\xcf\x11\xe0",
    ".ppt": b"\xd0\xcf\x11\xe0",
    ".docx": b"PK",
    ".xlsx": b"PK",
    ".pptx": b"PK",
}


def _is_html_response(content_type, filename):
    if "text/html" in content_type:
        ext = os.path.splitext(filename)[1].lower()
        return ext in MAGIC_BYTES
    return False


def _validate_file_magic(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".part":
        ext = os.path.splitext(os.path.splitext(filepath)[0])[1].lower()
    expected = MAGIC_BYTES.get(ext)
    if not expected:
        return True

    try:
        with open(filepath, "rb") as fh:
            header = fh.read(len(expected))
        return header.startswith(expected)
    except Exception:
        return True


# ---------------------------------------------------------------------------
# Metadata extraction
# ---------------------------------------------------------------------------

def extract_metadata(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".pptx": _extract_pptx,
        ".doc": _extract_ole,
        ".xls": _extract_ole,
        ".ppt": _extract_ole,
    }

    extractor = extractors.get(ext)
    if not extractor:
        logger.debug(f"No metadata extractor for {ext} files")
        return {}

    try:
        raw = extractor(filepath)
    except Exception as err:
        logger.warning(f"Failed to extract metadata from {os.path.basename(filepath)}: {err}")
        return {}

    cleaned = {}
    for key, value in raw.items():
        text = _clean_metadata_value(value)
        if text:
            cleaned[key] = text
    return cleaned


def _clean_metadata_value(value):
    """Make one metadata value safe to print and to write to a CSV.

    The target authored these strings, so they can carry terminal escape
    sequences, fake newlines meant to forge extra entries in the printed
    summary, or run to absurd lengths.
    """
    text = CONTROL_CHARS.sub("", str(value))
    text = " ".join(text.split())
    if len(text) > MAX_METADATA_LENGTH:
        text = text[:MAX_METADATA_LENGTH] + "..."
    return text


def _extract_pdf(filepath):
    from pypdf import PdfReader

    reader = PdfReader(filepath)
    info = reader.metadata
    if not info:
        return {}

    metadata = {}
    field_map = {
        "/Author": "Author",
        "/Creator": "Creator",
        "/Producer": "Producer",
        "/Title": "Title",
        "/Subject": "Subject",
        "/Keywords": "Keywords",
        "/CreationDate": "Created",
        "/ModDate": "Modified",
    }
    for pdf_key, our_key in field_map.items():
        value = info.get(pdf_key)
        if value:
            metadata[our_key] = str(value).strip()

    return metadata


def _extract_docx(filepath):
    from docx import Document

    doc = Document(filepath)
    props = doc.core_properties
    metadata = {}

    if props.author:
        metadata["Author"] = props.author
    if props.last_modified_by:
        metadata["LastModifiedBy"] = props.last_modified_by
    if props.title:
        metadata["Title"] = props.title
    if props.subject:
        metadata["Subject"] = props.subject
    if props.keywords:
        metadata["Keywords"] = props.keywords
    if props.created:
        metadata["Created"] = str(props.created)
    if props.modified:
        metadata["Modified"] = str(props.modified)

    _extract_ooxml_extended(filepath, metadata)

    return metadata


def _extract_xlsx(filepath):
    from openpyxl import load_workbook

    metadata = {}
    wb = load_workbook(filepath, read_only=True, data_only=True)
    try:
        props = wb.properties

        if props.creator:
            metadata["Author"] = props.creator
        if props.lastModifiedBy:
            metadata["LastModifiedBy"] = props.lastModifiedBy
        if props.title:
            metadata["Title"] = props.title
        if props.keywords:
            metadata["Keywords"] = props.keywords
        if props.created:
            metadata["Created"] = str(props.created)
        if props.modified:
            metadata["Modified"] = str(props.modified)
    finally:
        # read_only workbooks hold file handles open until closed.
        wb.close()

    _extract_ooxml_extended(filepath, metadata)

    return metadata


def _extract_pptx(filepath):
    from pptx import Presentation

    prs = Presentation(filepath)
    props = prs.core_properties
    metadata = {}

    if props.author:
        metadata["Author"] = props.author
    if props.last_modified_by:
        metadata["LastModifiedBy"] = props.last_modified_by
    if props.title:
        metadata["Title"] = props.title
    if props.subject:
        metadata["Subject"] = props.subject
    if props.keywords:
        metadata["Keywords"] = props.keywords
    if props.created:
        metadata["Created"] = str(props.created)
    if props.modified:
        metadata["Modified"] = str(props.modified)

    _extract_ooxml_extended(filepath, metadata)

    return metadata


OOXML_EXTENDED_FIELDS = {
    "manager": "Manager",
    "template": "Template",
    "application": "Application",
    "appversion": "AppVersion",
    "company": "Company",
}


def _extract_ooxml_extended(filepath, metadata):
    """Extract extended properties (Manager, Template, Application) from OOXML ZIP."""
    import zipfile

    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            try:
                info = zf.getinfo("docProps/app.xml")
            except KeyError:
                return
            # Deflate reaches ~1000:1, so an 80 KB entry can expand to 80 MB.
            if info.file_size > MAX_APP_XML_BYTES:
                logger.debug(f"Oversized app.xml ({info.file_size} bytes) in {filepath}")
                return
            app_xml = zf.read("docProps/app.xml")

        root = ElementTree.fromstring(app_xml)
        for element in root.iter():
            # Tags arrive namespaced as {...extended-properties}Manager
            tag = element.tag.rsplit("}", 1)[-1].lower()
            key = OOXML_EXTENDED_FIELDS.get(tag)
            if key and element.text and element.text.strip():
                metadata.setdefault(key, element.text.strip())
    except Exception as err:
        logger.debug(f"Could not read extended properties from {filepath}: {err}")


def _extract_ole(filepath):
    import olefile

    if not olefile.isOleFile(filepath):
        return {}

    ole = olefile.OleFileIO(filepath)
    metadata = {}

    try:
        meta = ole.get_metadata()

        fields = {
            "Author": meta.author,
            "LastModifiedBy": meta.last_saved_by,
            "Company": meta.company,
            "Title": meta.title,
            "Subject": meta.subject,
            "CreatingApplication": meta.creating_application,
            "Manager": meta.manager,
            "Keywords": meta.keywords,
            "Created": meta.create_time,
            "Modified": meta.last_saved_time,
        }
        for key, value in fields.items():
            if value:
                metadata[key] = _ole_text(value)
    finally:
        ole.close()

    return metadata


def _ole_text(value):
    """OLE property values are bytes for VT_LPSTR but str for VT_LPWSTR.

    Calling .decode() unconditionally raised AttributeError on the str form,
    and extract_metadata's except then threw away the whole file's metadata.
    """
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


# ---------------------------------------------------------------------------
# Output formatting
# ---------------------------------------------------------------------------

def format_metadata_summary(results):
    if not results:
        return "No metadata found."

    unique_users = set()
    unique_emails = set()
    unique_software = set()
    unique_companies = set()

    files_with_metadata = 0
    files_without_metadata = 0

    for entry in results:
        meta = entry["metadata"]
        if not meta:
            files_without_metadata += 1
            continue
        files_with_metadata += 1

        for field in ["Author", "LastModifiedBy", "Manager"]:
            if field in meta and meta[field]:
                unique_users.add(meta[field])

        for field in ["Creator", "Producer", "CreatingApplication"]:
            if field in meta and meta[field]:
                unique_software.add(meta[field])
        if "Application" in meta and meta["Application"]:
            app = meta["Application"]
            if "AppVersion" in meta:
                app = f"{app} {meta['AppVersion']}"
            unique_software.add(app)

        if "Company" in meta and meta["Company"]:
            unique_companies.add(meta["Company"])

        for value in meta.values():
            for email in EMAIL_REGEX.findall(str(value)):
                unique_emails.add(email.lower())

    lines = []
    lines.append("")
    lines.append("=" * 80)
    lines.append("METADATA SUMMARY")
    lines.append(f"  {files_with_metadata} file(s) with metadata, {files_without_metadata} without")
    lines.append("=" * 80)

    if unique_users:
        lines.append("")
        lines.append("NAMES / USERNAMES:")
        for user in sorted(unique_users, key=str.lower):
            lines.append(f"  - {user}")

    if unique_emails:
        lines.append("")
        lines.append("EMAILS:")
        for email in sorted(unique_emails):
            lines.append(f"  - {email}")

    if unique_companies:
        lines.append("")
        lines.append("COMPANIES:")
        for company in sorted(unique_companies, key=str.lower):
            lines.append(f"  - {company}")

    if unique_software:
        lines.append("")
        lines.append("SOFTWARE:")
        for sw in sorted(unique_software, key=str.lower):
            lines.append(f"  - {sw}")

    if not any([unique_users, unique_emails, unique_companies, unique_software]):
        lines.append("")
        lines.append("  No interesting metadata found in any files.")

    lines.append("")
    lines.append("=" * 80)

    return "\n".join(lines)


def export_csv(results, csv_path):
    if not results:
        logger.warning("No results to export.")
        return

    all_keys = set()
    for entry in results:
        all_keys.update(entry["metadata"].keys())
    all_keys = sorted(all_keys)

    fieldnames = ["Filename", "URL"] + all_keys

    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()

        for entry in results:
            row = {"Filename": _csv_safe(entry["file"]),
                   "URL": _csv_safe(entry.get("url", ""))}
            for key, value in entry["metadata"].items():
                row[key] = _csv_safe(value)
            writer.writerow(row)

    logger.info(f"Metadata exported to {csv_path}")


def _csv_safe(value):
    """Keep a spreadsheet from running a metadata value as a formula.

    An Author field of "=cmd|'/c calc'!A0" is a live payload when the CSV is
    opened in Excel. A leading apostrophe makes Excel treat it as text.
    """
    text = str(value)
    if text.startswith(CSV_FORMULA_PREFIXES):
        return f"'{text}"
    return text


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        prog="pyspyder",
        description=(
            "PySpyder - Spider a target domain for publicly available files, "
            "download them, and extract metadata."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python pyspyder.py -d targetdomain.com
  python pyspyder.py -d targetdomain.com -f pdf,docx
  python pyspyder.py -d targetdomain.com --depth 4
  python pyspyder.py -d targetdomain.com --csv results.csv
  python pyspyder.py --url-list urls.txt -o ./loot/
  python pyspyder.py --metadata-only ./loot/targetdomain.com-2024-03-25-143022
        """,
    )

    # Exactly one target mode. They used to be plain arguments, so passing two
    # was accepted and one of them silently ignored.
    target_group = parser.add_argument_group("target (exactly one required)")
    target_mode = target_group.add_mutually_exclusive_group()
    target_mode.add_argument(
        "-d", "--domain",
        help="Target domain to spider for files",
    )
    target_mode.add_argument(
        "--url-list",
        help="Path to a file containing URLs to process (one per line). Skips spidering.",
    )
    target_mode.add_argument(
        "--metadata-only",
        metavar="DIR",
        help="Extract metadata from files in an existing directory tree "
             "(skips spidering and downloading)",
    )

    parser.add_argument(
        "-f", "--file-types",
        default=",".join(DEFAULT_FILE_TYPES),
        help=f"Comma-separated file extensions to search for "
             f"(default: {','.join(DEFAULT_FILE_TYPES)})",
    )

    output_group = parser.add_argument_group("output")
    output_group.add_argument(
        "-o", "--output-dir",
        help="Parent directory for the timestamped run folder "
             "(default: current directory)",
    )
    output_group.add_argument(
        "--csv",
        metavar="FILE",
        help="Export all metadata to a CSV file",
    )

    spider_group = parser.add_argument_group("spider options")
    spider_group.add_argument(
        "--depth",
        type=int,
        default=2,
        help="Maximum crawl depth (default: 2)",
    )
    spider_group.add_argument(
        "--max-pages",
        type=int,
        default=500,
        help="Maximum pages to crawl (default: 500)",
    )
    spider_group.add_argument(
        "--ignore-robots",
        action="store_true",
        help="Ignore robots.txt when spidering and downloading",
    )

    spider_group.add_argument(
        "--download-only",
        action="store_true",
        help="Download files without extracting metadata",
    )

    request_group = parser.add_argument_group("request options")
    request_group.add_argument(
        "--delay",
        type=float,
        default=1.0,
        help="Delay in seconds between requests (default: 1.0)",
    )
    request_group.add_argument(
        "--user-agent",
        help="Custom User-Agent string",
    )
    request_group.add_argument(
        "--max-file-size",
        type=float,
        default=100.0,
        metavar="MB",
        help="Skip files larger than this many megabytes (default: 100, 0 for no limit)",
    )

    parser.add_argument(
        "-v", "--verbose",
        action="store_true",
        help="Enable verbose/debug output",
    )
    parser.add_argument(
        "--version",
        action="version",
        version=f"%(prog)s {__version__}",
    )

    args = parser.parse_args(argv)

    if not args.domain and not args.url_list and not args.metadata_only:
        parser.error("Either --domain (-d), --url-list, or --metadata-only is required.")

    if args.csv and args.download_only:
        parser.error("--csv cannot be used with --download-only "
                     "(no metadata is collected to export).")

    if args.max_file_size < 0:
        parser.error("--max-file-size cannot be negative.")

    if args.delay < 0:
        parser.error("--delay cannot be negative.")

    if args.depth < 0:
        parser.error("--depth cannot be negative.")

    if args.max_pages < 1:
        parser.error("--max-pages must be at least 1.")

    return args


def collect_metadata(paths_and_urls):
    results = []
    for name, url, filepath in paths_and_urls:
        results.append({
            "file": name,
            "url": url,
            "metadata": extract_metadata(filepath),
        })
    return results


def dedupe_urls(urls):
    """One URL per resource, sorted.

    Deduplicates on the canonical key but returns each URL exactly as it was
    discovered, since that is what gets requested. Prefers https when the same
    resource turned up under both schemes, and iterates in sorted order so the
    survivor is the same on every run.
    """
    by_key = {}
    for url in sorted(urls):
        key = url_key(url)
        incumbent = by_key.get(key)
        if incumbent is None or (url.startswith("https://")
                                 and not incumbent.startswith("https://")):
            by_key[key] = url
    return sorted(by_key.values())


def find_local_files(target_dir, file_types):
    """Walk a directory tree and return (display_name, path) for matching files."""
    wanted = {ft.lower().lstrip(".") for ft in file_types}
    found = []
    for dirpath, _, filenames in os.walk(target_dir):
        for filename in sorted(filenames):
            ext = os.path.splitext(filename)[1].lstrip(".").lower()
            if ext not in wanted:
                continue
            full = os.path.join(dirpath, filename)
            found.append((os.path.relpath(full, target_dir), full))
    return sorted(found)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    log = setup_logging(args.verbose)

    file_types = normalize_file_types(args.file_types)
    if not file_types:
        log.error("No file types given to --file-types.")
        sys.exit(1)

    max_bytes = int(args.max_file_size * 1024 * 1024)

    # Metadata-only mode: skip spidering and downloading
    if args.metadata_only:
        target_dir = args.metadata_only
        if not os.path.isdir(target_dir):
            log.error(f"Directory not found: {target_dir}")
            sys.exit(1)

        log.info(f"Extracting metadata from files in: {target_dir}")
        local_files = find_local_files(target_dir, file_types)
        results = collect_metadata([(name, "", path) for name, path in local_files])

        print(f"\nProcessed {len(results)} file(s).\n")
        print(format_metadata_summary(results))

        if args.csv:
            export_csv(results, args.csv)
        return

    session = get_session(args.user_agent)

    # Phase 1: Discover file URLs
    all_urls = set()
    domain_name = "url-list"

    if args.url_list:
        log.info(f"Loading URLs from {args.url_list}")
        try:
            with open(args.url_list, "r") as fh:
                for line in fh:
                    line = line.strip()
                    if not line or line.startswith("#"):
                        continue
                    if "://" not in line:
                        line = f"https://{line}"
                        log.debug(f"Assuming https:// for {line}")
                    if not line.lower().startswith(("http://", "https://")):
                        log.warning(f"Skipping unsupported URL: {line}")
                        continue
                    all_urls.add(line)
            log.info(f"Loaded {len(all_urls)} URL(s) from file")
        except FileNotFoundError:
            log.error(f"URL list file not found: {args.url_list}")
            sys.exit(1)
    else:
        domain_name, discarded_path = normalize_domain(args.domain)
        if not domain_name:
            log.error(f"Could not read a hostname from --domain {args.domain!r}")
            sys.exit(1)
        if discarded_path:
            log.warning(f"Ignoring path '/{discarded_path}' - the crawl starts at "
                        f"the site root of {domain_name}")

        try:
            all_urls.update(spider_domain(
                domain_name, file_types, session,
                max_depth=args.depth,
                delay=args.delay,
                ignore_robots=args.ignore_robots,
                max_pages=args.max_pages,
            ))
        except Exception as err:
            log.error(f"Spider failed: {err}")
            sys.exit(1)

    all_urls = dedupe_urls(all_urls)
    print(f"\nFound {len(all_urls)} unique file(s) to download.\n")

    if not all_urls:
        log.info("No files found. Exiting.")
        sys.exit(0)

    # Phase 2: Download
    run_dir = generate_output_dir(domain_name)
    output_dir = os.path.join(args.output_dir, run_dir) if args.output_dir else run_dir

    log.info(f"Downloading files to: {output_dir}")
    scope_hosts = {bare_host(domain_name)} if args.domain else None
    downloaded = download_files(all_urls, output_dir, session, delay=args.delay,
                                max_bytes=max_bytes, scope_hosts=scope_hosts)
    print(f"\nDownloaded {len(downloaded)} of {len(all_urls)} file(s).\n")

    if not downloaded:
        log.info("No files downloaded. Exiting.")
        sys.exit(0)

    if args.download_only:
        print(f"\nFiles saved to: {os.path.abspath(output_dir)}")
        return

    # Phase 3: Extract metadata
    results = collect_metadata(
        [(os.path.basename(path), url, path) for url, path in downloaded]
    )

    print(format_metadata_summary(results))

    if args.csv:
        export_csv(results, args.csv)

    print(f"\nFiles saved to: {os.path.abspath(output_dir)}")


if __name__ == "__main__":
    main()
